"""
Objective-driven analysis over one or more uploaded files.

Workflow the page follows:

1. **Review** - every uploaded file is checked with plain rules (`data_health`):
   header problems, blank columns, unreadable or mixed-format dates, amounts
   stored as text, stray spaces, duplicate IDs, TOTAL rows. `find_links` looks
   for columns that connect files (e.g. a contract number in both). Claude then
   writes "what I see" from those findings and a column profile - never from the
   full data - and recommends formatting fixes.
2. **Objectives, one after another** - each new objective refines the last.
   Claude returns the *complete* plan for the whole conversation so far, built
   only from the operations in `OPS`; the final product is therefore always
   recalculated from the original files and is reproducible.
3. **Calculation** - `execute_plan` runs the approved steps in plain pandas. No AI
   touches a number; the same plan on the same files gives the same answer.
4. **Summary** - a short low-effort call writes the opening from the result
   table only.

James is the only user and authorized AI use (2026-10-20). Classification stays
rules-only; ANTHROPIC_API_KEY is needed only for the review narrative, planning
and summary.
"""
from __future__ import annotations

import io
import json
import re
from datetime import date

import pandas as pd

import spend_report_JHK3 as sr

MODEL = "claude-opus-5-5"
FALLBACK_BETA = "server-side-fallback-2026-07-01"
SAMPLE_VALUES = 6
MAX_PROFILE_COLS = 80
RESULT_ROWS_FOR_SUMMARY = 40
LINK_SAMPLE = 5000

# --------------------------------------------------------------------------
# Operations a plan may use
# --------------------------------------------------------------------------
OPS = {
    "start_from": "Begin with the file named in `value`. Use as the first step whenever more than one "
                  "file is loaded; otherwise the first file is used.",
    "join": "Bring in columns from the file named in `value`, matching `columns` (keys in the current "
            "table) to `values` (keys in that file, same order). `operator` is 'left' (keep every current "
            "row) or 'inner' (keep only rows that match). Keys are compared ignoring case, spaces and "
            "a trailing '.0'.",
    "append": "Stack the rows of the file named in `value` under the current table (columns matched by "
              "name). Adds a 'Source file' column.",
    "filter": "Keep rows where `column` meets `operator` and `value`/`values`. Operators: equals, "
              "not_equals, contains, not_contains, in, not_in, gt, gte, lt, lte, between, is_blank, "
              "not_blank. Text comparisons ignore case. Numbers and dates (YYYY-MM-DD) are parsed "
              "automatically.",
    "days_until": "Add column `as_name` = whole days from today to the date in `column` (negative = past).",
    "date_part": "Add column `as_name` = the year, quarter or month of the date in `column`; put "
                 "'year', 'quarter' or 'month' in `operator`.",
    "calculate": "Add column `as_name` = `column` combined with `value` using `operator`: add, subtract, "
                 "multiply, divide, or percent_of (column ÷ value × 100). `value` is another column's name "
                 "or a number.",
    "bucket": "Add column `as_name` sorting the number in `column` into ranges. `edges` ascending; "
              "`labels` has one fewer item than edges.",
    "group_summarize": "Group by `columns` and compute `metrics` (source column, agg: sum, count, mean, "
                       "min, max, nunique, and output name).",
    "add_share": "Add column `as_name` = each row's share (%) of the total of `column`.",
    "sort": "Sort by `column` (largest/latest first when `descending` is true).",
    "top_n": "Keep the first `n` rows.",
    "select_columns": "Keep only `columns`, in that order.",
    "rename_columns": "Rename: `columns` are old names, `values` the new names, in order.",
}

_STEP_SCHEMA = {
    "type": "object",
    "properties": {
        "op": {"type": "string", "enum": list(OPS)},
        "column": {"type": "string"},
        "columns": {"type": "array", "items": {"type": "string"}},
        "operator": {"type": "string"},
        "value": {"type": "string"},
        "values": {"type": "array", "items": {"type": "string"}},
        "as_name": {"type": "string"},
        "edges": {"type": "array", "items": {"type": "number"}},
        "labels": {"type": "array", "items": {"type": "string"}},
        "metrics": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "column": {"type": "string"},
                    "agg": {"type": "string", "enum": ["sum", "count", "mean", "min", "max", "nunique"]},
                    "as_name": {"type": "string"},
                },
                "required": ["column", "agg", "as_name"],
                "additionalProperties": False,
            },
        },
        "descending": {"type": "boolean"},
        "n": {"type": "integer"},
        "why": {"type": "string"},
    },
    "required": ["op", "column", "columns", "operator", "value", "values", "as_name", "edges",
                 "labels", "metrics", "descending", "n", "why"],
    "additionalProperties": False,
}

PLAN_SCHEMA = {
    "type": "object",
    "properties": {
        "restated_objective": {"type": "string"},
        "assumptions": {"type": "array", "items": {"type": "string"}},
        "questions": {"type": "array", "items": {"type": "string"}},
        "cannot_do": {"type": "array", "items": {"type": "string"}},
        "steps": {"type": "array", "items": _STEP_SCHEMA},
        "result_title": {"type": "string"},
    },
    "required": ["restated_objective", "assumptions", "questions", "cannot_do", "steps", "result_title"],
    "additionalProperties": False,
}

OVERVIEW_SCHEMA = {
    "type": "object",
    "properties": {
        "summary": {"type": "string"},
        "files": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "name": {"type": "string"},
                    "appears_to_be": {"type": "string"},
                    "key_columns": {"type": "array", "items": {"type": "string"}},
                    "useful_for": {"type": "array", "items": {"type": "string"}},
                },
                "required": ["name", "appears_to_be", "key_columns", "useful_for"],
                "additionalProperties": False,
            },
        },
        "relationships": {"type": "array", "items": {"type": "string"}},
        "can_answer": {"type": "array", "items": {"type": "string"}},
        "formatting_recommendations": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "priority": {"type": "string", "enum": ["High", "Medium", "Low"]},
                    "file": {"type": "string"},
                    "column": {"type": "string"},
                    "issue": {"type": "string"},
                    "fix": {"type": "string"},
                },
                "required": ["priority", "file", "column", "issue", "fix"],
                "additionalProperties": False,
            },
        },
    },
    "required": ["summary", "files", "relationships", "can_answer", "formatting_recommendations"],
    "additionalProperties": False,
}

PLANNER_SYSTEM = """You plan analyses of procurement data for a public-sector contracts administrator \
(a city water department). You never calculate anything yourself: you write a plan that a separate, \
deterministic program runs, and the administrator must be able to defend every number to leadership.

Rules:
- Use only the operations listed. Use only file names and column names from the data profile, or \
names created by an earlier step's as_name / metric as_name / a join.
- When more than one file is loaded, the first step is start_from. Join only on columns the profile \
or the links list shows can connect; say in assumptions which keys you joined on.
- Every step must fill every field. Unused text fields are "", unused lists are [], unused numbers \
are 0, unused booleans are false. Give each step a short plain-English `why`.
- The administrator builds the result over several objectives. You receive the earlier objectives \
in order and the current plan. Return the COMPLETE plan for everything asked so far, starting from \
the source files - not just the new steps. Later instructions override earlier ones.
- Restate the whole request in one sentence as you now understand it.
- List every assumption (which column holds the end date, what "this year" means, what counts as a \
chemical). Prefer columns over guesses.
- If the request is ambiguous in a way that would change the answer, put the question in \
`questions` and still give your best plan.
- If something needs an operation, column or file that isn't available, say so plainly in \
`cannot_do` and plan what is possible. Never invent data.
- End with a table a leader can read: few columns, clear names, most important rows first.
- Columns Business_Category, NIGP_Class_3digit, NIGP_Item_5digit, NAICS_Code and NAICS_Title, when \
present, come from the mapper's rules-based classification of a description column."""

OVERVIEW_SYSTEM = """You review uploaded procurement files for a public-sector contracts administrator \
before any analysis. You see a profile of each file (columns, kinds, sample values, counts), the \
results of automatic data checks, and possible links between files. You do not see the full data.

Write what you see, plainly:
- summary: two or three sentences on what this set of files is and what it can support.
- files: for each file, what it appears to be (e.g. a contract list, a payment history), its key \
columns, and what it is useful for.
- relationships: how the files connect (which columns match, roughly how well, per the links list).
- can_answer: five to eight concrete questions these files can answer for a contracts administrator \
(expirations, spend by vendor, contracts without payments, etc.). Only questions the columns support.
- formatting_recommendations: what to fix in the source files before relying on results, most \
important first. Ground every item in the automatic checks or the profile; include the automatic \
checks' items that matter, merged and reworded for a non-technical reader, and add any you can see \
in the profile (for example a date column the checks missed, or a column that mixes two kinds of \
value). Priority High = would change an answer; Medium = may cause missed matches or confusion; \
Low = tidiness. The fix says exactly what to do in Excel.
Do not invent columns or problems you cannot see."""

SUMMARY_SYSTEM = """You write the opening of a briefing for a public-sector contracts administrator. \
You are given the request, the steps that were run, and the result table. Lead with the direct answer \
in one or two sentences, using numbers from the table exactly as they appear. Then give at most three \
short supporting points. Do not add numbers that are not in the table, do not speculate about causes, \
and mention any caveat listed. Plain language, no headings, no jargon."""


# --------------------------------------------------------------------------
# Profiles and checks (deterministic)
# --------------------------------------------------------------------------
_KEY_HINTS = ("number", "#", " id", "id ", "po", "contract", "spec", "vendor id", "invoice", "voucher")
_DATE_HINTS = ("date", "start", "end", "expir", "award")
_DATE_PATTERNS = {
    "YYYY-MM-DD": r"^\d{4}-\d{1,2}-\d{1,2}",
    "MM/DD/YYYY": r"^\d{1,2}/\d{1,2}/\d{2,4}",
    "text month (e.g. Jan 5, 2026)": r"^[A-Za-z]{3,9}\.? \d{1,2},? \d{2,4}",
    "DD-Mon-YY": r"^\d{1,2}-[A-Za-z]{3}-\d{2,4}",
}


def _kind(s: pd.Series) -> str:
    if pd.api.types.is_datetime64_any_dtype(s):
        return "date"
    if pd.api.types.is_numeric_dtype(s):
        return "number"
    if sr._date_frac(s) >= 0.7:
        return "date (stored as text)"
    if sr._num_frac(s) >= 0.8:
        return "number (stored as text)"
    return "text"


def profile_file(name: str, df: pd.DataFrame) -> dict:
    cols = []
    for c in list(df.columns)[:MAX_PROFILE_COLS]:
        s = df[c]
        nn = s.dropna()
        nn = nn[nn.astype(str).str.strip() != ""]
        samples = [str(v)[:60] for v in nn.astype(str).drop_duplicates().head(SAMPLE_VALUES)]
        cols.append({"name": str(c), "kind": _kind(s), "non_blank": int(len(nn)),
                     "distinct": int(nn.nunique()), "examples": samples})
    return {"file": name, "rows": int(len(df)), "columns": cols,
            "columns_not_shown": max(0, len(df.columns) - MAX_PROFILE_COLS)}


def profile_files(files: dict) -> dict:
    return {"today": date.today().isoformat(),
            "files": [profile_file(n, d) for n, d in files.items()]}


def profile_for_planner(df: pd.DataFrame) -> dict:
    """Single-file profile (kept for callers that pass one DataFrame)."""
    return profile_files({"data": df})


def _norm_key(s: pd.Series) -> pd.Series:
    return (s.astype(str).str.strip().str.upper()
            .str.replace(r"\.0$", "", regex=True).str.replace(r"\s+", " ", regex=True))


def _is_keyish(name: str, s: pd.Series) -> bool:
    low = f" {str(name).lower()} "
    nn = s.dropna()
    card = nn.nunique() / len(nn) if len(nn) else 0
    return any(h in low for h in _KEY_HINTS) or (card >= 0.9 and len(nn) >= 20
                                                  and nn.astype(str).str.len().mean() < 25)


def data_health(name: str, df: pd.DataFrame) -> list:
    """Plain-rule checks. Each finding names the file, column, issue, fix, priority, rows."""
    out = []

    def add(col, issue, fix, priority, rows=0):
        out.append({"file": name, "column": col, "issue": issue, "fix": fix,
                    "priority": priority, "rows_affected": int(rows)})

    unnamed = [c for c in df.columns if str(c).lower().startswith("unnamed")]
    if unnamed:
        add(", ".join(map(str, unnamed[:5])),
            f"{len(unnamed)} column(s) have no header name.",
            "Make row 1 hold the column names (delete any title rows above it) and delete blank columns.",
            "High" if len(unnamed) > len(df.columns) / 3 else "Low")
    dup_rows = int(df.duplicated().sum())
    if dup_rows:
        add("(whole row)", f"{dup_rows:,} rows are exact duplicates of another row.",
            "Check whether they are true repeats; delete them if so, or totals will double-count.",
            "High", dup_rows)
    first_text = next((c for c in df.columns if _is_text(df[c])), None)
    if first_text is not None:
        tot = df[first_text].astype(str).str.strip().str.fullmatch(r"(?i)(grand )?totals?:?")
        if tot.any():
            add(str(first_text), "The file appears to contain a TOTAL row.",
                "Delete TOTAL / subtotal rows so they aren't counted as records.", "High", int(tot.sum()))

    for c in df.columns:
        s = df[c]
        n = len(s)
        if not n:
            continue
        blank = s.isna() | (s.astype(str).str.strip() == "")
        nb = int(blank.sum())
        name_low = str(c).lower()
        if nb == n:
            add(str(c), "The column is completely empty.", "Delete the column.", "Low", nb)
            continue
        if nb / n > 0.5:
            add(str(c), f"{nb:,} of {n:,} rows ({nb / n:.0%}) are blank.",
                "Fill in the missing values if this column matters for your questions.", "Medium", nb)
        vals = s[~blank]
        if _is_text(vals):
            sv = vals.astype(str)
            spaced = int((sv != sv.str.strip()).sum())
            if spaced and spaced / len(sv) > 0.01:
                add(str(c), f"{spaced:,} values have extra spaces before or after the text.",
                    "Use Excel TRIM (or Find & Replace) so values match across files.", "Medium", spaced)
            # dates
            datef = sr._date_frac(vals)
            if any(h in name_low for h in _DATE_HINTS) or datef >= 0.3:
                parsed = _as_date(vals)
                bad = int(parsed.isna().sum())
                if datef >= 0.3 and bad:
                    eg = ", ".join(f"'{v}'" for v in sv[parsed.isna()].drop_duplicates().head(3))
                    add(str(c), f"{bad:,} values aren't readable dates (e.g. {eg}).",
                        "Enter a real date or leave the cell blank; use one format such as 2026-12-31.",
                        "High", bad)
                shares = {k: sv.str.match(p).mean() for k, p in _DATE_PATTERNS.items()}
                used = [k for k, v in shares.items() if v >= 0.05]
                if len(used) > 1:
                    add(str(c), f"Dates are written in {len(used)} different formats ({', '.join(used)}).",
                        "Format the whole column as one date format (e.g. YYYY-MM-DD) so dates sort and "
                        "compare correctly.", "Medium", len(sv))
            # amounts stored as text
            numf = sr._num_frac(vals)
            if numf >= 0.8 and sv.str.contains(r"[$,]|up\s*to|nte", case=False, regex=True).mean() > 0.2:
                add(str(c), "Amounts are stored as text (with $, commas or 'Up To').",
                    "The tool reads these, but for Excel math store them as plain numbers (Number format).",
                    "Low", len(sv))
            elif 0.3 <= numf < 0.8 and any(h in name_low for h in ("amount", "value", "cost", "paid", "total")):
                bad = int(sr.clean_amount(vals).isna().sum())
                add(str(c), f"{bad:,} values in this amount column aren't readable numbers.",
                    "Replace notes such as 'TBD' with a number or leave blank; keep notes in another column.",
                    "High", bad)
        # keys
        if _is_keyish(c, s) and len(vals) >= 2:
            kv = _norm_key(vals)
            d = int(kv.duplicated().sum())
            if d and d / len(kv) < 0.5:
                add(str(c), f"{d:,} repeated values in what looks like an ID column.",
                    "Check whether these are revisions or duplicate rows; keep one row per ID if the "
                    "file should list each contract once.", "Medium", d)
            if pd.api.types.is_float_dtype(s):
                add(str(c), "IDs are stored as decimal numbers (e.g. 296724.0), which can drop leading "
                    "zeros.", "Format the column as Text before typing or pasting IDs.", "Medium", len(vals))
    order = {"High": 0, "Medium": 1, "Low": 2}
    return sorted(out, key=lambda r: (order[r["priority"]], -r["rows_affected"]))


def find_links(files: dict) -> list:
    """Columns that connect two files: share of one file's non-blank values found in the other."""
    names = list(files)
    keyish = {n: [c for c in files[n].columns if _is_keyish(c, files[n][c])] for n in names}
    links = []
    for i, a in enumerate(names):
        for b in names[i + 1:]:
            for ca in keyish[a]:
                va = set(_norm_key(files[a][ca].dropna().head(LINK_SAMPLE)))
                va.discard("")
                if not va:
                    continue
                for cb in keyish[b]:
                    vb = set(_norm_key(files[b][cb].dropna().head(LINK_SAMPLE)))
                    vb.discard("")
                    if not vb:
                        continue
                    share = len(va & vb) / len(va)
                    if share >= 0.2:
                        links.append({"file_a": a, "column_a": str(ca), "file_b": b, "column_b": str(cb),
                                      "match_share": round(share, 2),
                                      "note": f"{share:.0%} of '{ca}' values in {a} also appear in "
                                              f"'{cb}' in {b}"})
    return sorted(links, key=lambda r: -r["match_share"])


# --------------------------------------------------------------------------
# Claude calls
# --------------------------------------------------------------------------
class AIUnavailable(RuntimeError):
    pass


def _client(api_key: str):
    if not api_key:
        raise AIUnavailable("No Anthropic API key is set. Add ANTHROPIC_API_KEY to the app's secrets.")
    import anthropic
    return anthropic.Anthropic(api_key=api_key)


def _text_of(response) -> str:
    if response.stop_reason == "refusal":
        raise AIUnavailable("The request was declined. Rephrase it and try again.")
    if response.stop_reason == "max_tokens":
        raise AIUnavailable("The response was cut off. Try a narrower request.")
    return next((b.text for b in response.content if b.type == "text"), "")


def _ask_json(api_key, system, user, schema, effort="high"):
    resp = _client(api_key).beta.messages.create(
        model=MODEL,
        max_tokens=16000,
        betas=[FALLBACK_BETA],
        fallbacks="default",
        output_config={"effort": effort, "format": {"type": "json_schema", "schema": schema}},
        system=system,
        messages=[{"role": "user", "content": user}],
    )
    return json.loads(_text_of(resp))


def overview_files(profile: dict, health: list, links: list, api_key: str) -> dict:
    user = (f"Data profile (JSON):\n{json.dumps(profile, indent=1)}\n\n"
            f"Automatic checks (JSON):\n{json.dumps(health, indent=1)}\n\n"
            f"Possible links between files (JSON):\n{json.dumps(links, indent=1)}")
    return _ask_json(api_key, OVERVIEW_SYSTEM, user, OVERVIEW_SCHEMA, effort="medium")


def plan_objective(objective: str, profile: dict, api_key: str, answers: str = "",
                   history: list | None = None, current_plan: dict | None = None,
                   links: list | None = None) -> dict:
    """Plan for the whole conversation so far. `history` = earlier objectives in order."""
    ops = "\n".join(f"- {k}: {v}" for k, v in OPS.items())
    parts = [f"Operations available:\n{ops}",
             f"Data profile (JSON):\n{json.dumps(profile, indent=1)}"]
    if links:
        parts.append(f"Columns that connect files (JSON):\n{json.dumps(links, indent=1)}")
    if history:
        parts.append("Earlier objectives, in order:\n" + "\n".join(f"{i}. {h}" for i, h in enumerate(history, 1)))
    if current_plan:
        parts.append(f"Current plan (JSON):\n{json.dumps(current_plan, indent=1)}")
    parts.append(f"New objective:\n{objective.strip()}")
    if answers.strip():
        parts.append(f"The administrator answered your earlier questions:\n{answers.strip()}")
    return _ask_json(api_key, PLANNER_SYSTEM, "\n\n".join(parts), PLAN_SCHEMA, effort="high")


def summarize_result(request: str, plan: dict, result: pd.DataFrame, caveats: list, api_key: str) -> str:
    table = result.head(RESULT_ROWS_FOR_SUMMARY).to_csv(index=False)
    more = (f"\n(The full result has {len(result):,} rows; the first {RESULT_ROWS_FOR_SUMMARY} are shown.)"
            if len(result) > RESULT_ROWS_FOR_SUMMARY else "")
    user = (f"Request: {request}\n\nSteps run:\n" + "\n".join(f"- {s}" for s in describe_steps(plan))
            + "\n\nCaveats:\n" + ("\n".join(f"- {c}" for c in caveats) or "- none")
            + f"\n\nResult table (CSV):\n{table}{more}")
    resp = _client(api_key).beta.messages.create(
        model=MODEL, max_tokens=4000, betas=[FALLBACK_BETA], fallbacks="default",
        output_config={"effort": "low"}, system=SUMMARY_SYSTEM,
        messages=[{"role": "user", "content": user}],
    )
    return _text_of(resp).strip()


# --------------------------------------------------------------------------
# Plain-English description of a plan
# --------------------------------------------------------------------------
_OP_WORDS = {"equals": "is", "not_equals": "is not", "contains": "contains", "not_contains": "does not contain",
             "in": "is one of", "not_in": "is not one of", "gt": "is more than", "gte": "is at least",
             "lt": "is less than", "lte": "is at most", "between": "is between",
             "is_blank": "is blank", "not_blank": "is not blank"}
_CALC_WORDS = {"add": "plus", "subtract": "minus", "multiply": "times", "divide": "divided by",
               "percent_of": "as a % of"}


def describe_step(s: dict) -> str:
    op = s.get("op")
    if op == "start_from":
        return f"Start from **{s.get('value')}**"
    if op == "join":
        keys = ", ".join(f"{a} = {b}" for a, b in zip(s.get("columns") or [], s.get("values") or []))
        how = "keeping every row" if s.get("operator", "left") == "left" else "keeping only matches"
        return f"Bring in **{s.get('value')}**, matching {keys} ({how})"
    if op == "append":
        return f"Add the rows of **{s.get('value')}** underneath"
    if op == "filter":
        oper = s.get("operator", "")
        val = (" and ".join(s.get("values") or []) if oper == "between"
               else ", ".join(s.get("values") or []) or s.get("value", ""))
        cond = _OP_WORDS.get(oper, oper)
        return f"Keep rows where **{s['column']}** {cond}" + (f" **{val}**" if val and "blank" not in oper else "")
    if op == "days_until":
        return f"Count days from today to **{s['column']}** → **{s['as_name']}**"
    if op == "date_part":
        return f"Take the {s.get('operator')} of **{s['column']}** → **{s['as_name']}**"
    if op == "calculate":
        return (f"Calculate **{s['column']}** {_CALC_WORDS.get(s.get('operator'), s.get('operator'))} "
                f"**{s.get('value')}** → **{s['as_name']}**")
    if op == "bucket":
        return f"Sort **{s['column']}** into ranges {', '.join(s.get('labels') or [])} → **{s['as_name']}**"
    if op == "group_summarize":
        ms = "; ".join(f"{m['agg']} of {m['column']} as **{m['as_name']}**" for m in s.get("metrics") or [])
        return f"Group by **{', '.join(s.get('columns') or [])}** and compute {ms}"
    if op == "add_share":
        return f"Add each row's % of total **{s['column']}** → **{s['as_name']}**"
    if op == "sort":
        return f"Sort by **{s['column']}** ({'largest first' if s.get('descending') else 'smallest first'})"
    if op == "top_n":
        return f"Keep the top **{s.get('n')}** rows"
    if op == "select_columns":
        return f"Show columns **{', '.join(s.get('columns') or [])}**"
    if op == "rename_columns":
        return "Rename " + ", ".join(f"{a} → {b}" for a, b in zip(s.get("columns") or [], s.get("values") or []))
    return f"Unknown step {op}"


def describe_steps(plan: dict) -> list:
    return [describe_step(s) for s in plan.get("steps", [])]


# --------------------------------------------------------------------------
# Deterministic execution
# --------------------------------------------------------------------------
class PlanError(ValueError):
    pass


def _rows(n: int, past: bool = False) -> str:
    verb = ("had" if past else "has") if n == 1 else ("had" if past else "have")
    return f"{n:,} row{'' if n == 1 else 's'} {verb}"


def _need(work, col, step_no):
    if col not in work.columns:
        raise PlanError(f"Step {step_no}: there is no column named '{col}'. "
                        f"Available: {', '.join(map(str, work.columns))}")


def _need_file(files, name, step_no):
    if name not in files:
        raise PlanError(f"Step {step_no}: there is no file named '{name}'. Loaded: {', '.join(files)}")
    return files[name]


def _as_number(series):
    return sr.clean_amount(series)


def _as_date(series):
    if pd.api.types.is_datetime64_any_dtype(series):
        return series
    try:
        return pd.to_datetime(series, errors="coerce", format="mixed")
    except (ValueError, TypeError):
        return pd.to_datetime(series, errors="coerce")


def _is_text(series) -> bool:
    """True for text columns under both pandas 2 (object) and pandas 3 (str dtype)."""
    return not (pd.api.types.is_numeric_dtype(series) or pd.api.types.is_datetime64_any_dtype(series)
                or pd.api.types.is_bool_dtype(series))


def _is_dateish(series):
    return pd.api.types.is_datetime64_any_dtype(series) or sr._date_frac(series) >= 0.7


def _looks_date(v):
    v = str(v)
    if not re.match(r"^\d{4}-\d{1,2}-\d{1,2}", v):
        return False
    try:
        pd.Timestamp(v)
        return True
    except (ValueError, TypeError):
        return False


def _looks_numeric_or_date(v):
    return _looks_date(v) or pd.notna(sr.clean_amount(pd.Series([v])).iloc[0])


def _compare(series, operator, value, values):
    if operator == "is_blank":
        return series.isna() | (series.astype(str).str.strip() == "")
    if operator == "not_blank":
        return ~(series.isna() | (series.astype(str).str.strip() == ""))
    text = series.astype(str).str.strip().str.lower()
    if operator in ("contains", "not_contains"):
        m = text.str.contains(str(value).strip().lower(), regex=False, na=False)
        return m if operator == "contains" else ~m
    if operator in ("in", "not_in"):
        opts = {str(v).strip().lower() for v in (values or [value])}
        m = text.isin(opts)
        return m if operator == "in" else ~m
    if operator in ("equals", "not_equals") and not _looks_numeric_or_date(value):
        m = text == str(value).strip().lower()
        return m if operator == "equals" else ~m
    probe = (values or [""])[0] if operator == "between" else value
    if _is_dateish(series) and _looks_date(probe):
        left, conv = _as_date(series), (lambda v: pd.Timestamp(v))
    else:
        left, conv = _as_number(series), (lambda v: float(sr.clean_amount(pd.Series([v])).iloc[0]))
    if operator == "between":
        if not values or len(values) != 2:
            raise PlanError("A 'between' filter needs exactly two values.")
        return (left >= conv(values[0])) & (left <= conv(values[1]))
    v = conv(value)
    return {"equals": left == v, "not_equals": left != v, "gt": left > v, "gte": left >= v,
            "lt": left < v, "lte": left <= v}[operator]


def execute_plan(data, plan: dict, today: date | None = None):
    """Run the plan over `data` (one DataFrame, or {file name: DataFrame}).
    Returns (result DataFrame, list of caveats)."""
    files = data if isinstance(data, dict) else {"data": data}
    if not files:
        raise PlanError("No files are loaded.")
    today = pd.Timestamp(today or date.today())
    base_name = next(iter(files))
    work = files[base_name].copy()
    caveats = []
    for i, s in enumerate(plan.get("steps", []), 1):
        op = s.get("op")
        before = len(work)
        if op == "start_from":
            work = _need_file(files, s.get("value", ""), i).copy()
            base_name = s["value"]
        elif op == "join":
            other = _need_file(files, s.get("value", ""), i)
            lk, rk = list(s.get("columns") or []), list(s.get("values") or [])
            if not lk or len(lk) != len(rk):
                raise PlanError(f"Step {i}: a join needs the same number of keys on both sides.")
            for c in lk:
                _need(work, c, i)
            for c in rk:
                _need(other, c, i)
            how = s.get("operator") if s.get("operator") in ("left", "inner") else "left"
            left, right = work.copy(), other.copy()
            tmp = [f"__k{j}" for j in range(len(lk))]
            for t, a, b in zip(tmp, lk, rk):
                left[t], right[t] = _norm_key(left[a]), _norm_key(right[b])
            right = right.drop(columns=rk)
            clash = {c: f"{c} ({s['value']})" for c in right.columns if c in left.columns and c not in tmp}
            right = right.rename(columns=clash)
            matched = left[tmp].apply(tuple, axis=1).isin(set(right[tmp].apply(tuple, axis=1)))
            dupes = int(right.duplicated(subset=tmp).sum())
            work = left.merge(right, on=tmp, how=how).drop(columns=tmp)
            caveats.append(f"Step {i}: {int(matched.sum()):,} of {len(left):,} rows found a match in "
                           f"{s['value']}.")
            if dupes:
                caveats.append(f"Step {i}: {s['value']} has more than one row for some keys (normal for a "
                               f"payment history), so matched rows repeat until a later step totals them.")
        elif op == "append":
            other = _need_file(files, s.get("value", ""), i).copy()
            if "Source file" not in work.columns:
                work = work.assign(**{"Source file": base_name})
            other["Source file"] = s["value"]
            work = pd.concat([work, other], ignore_index=True, sort=False)
        elif op == "filter":
            _need(work, s["column"], i)
            if s.get("operator") not in _OP_WORDS:
                raise PlanError(f"Step {i}: unknown filter operator '{s.get('operator')}'.")
            col = work[s["column"]]
            work = work[_compare(col, s["operator"], s.get("value", ""), s.get("values"))]
            if s["operator"] in ("gt", "gte", "lt", "lte", "between"):
                parsed = _as_date(col) if _is_dateish(col) else _as_number(col)
                blanks = int(parsed.isna().sum())
                if blanks:
                    caveats.append(f"{_rows(blanks, past=True)} no usable value in '{s['column']}' and "
                                   f"{'was' if blanks == 1 else 'were'} left out at step {i}.")
        elif op == "days_until":
            _need(work, s["column"], i)
            d = _as_date(work[s["column"]])
            work[s["as_name"]] = (d - today).dt.days
            if d.isna().any():
                caveats.append(f"{_rows(int(d.isna().sum()))} no usable date in '{s['column']}'.")
        elif op == "date_part":
            _need(work, s["column"], i)
            d = _as_date(work[s["column"]])
            part = s.get("operator", "year")
            if part == "year":
                work[s["as_name"]] = d.dt.year.astype("Int64")
            elif part == "quarter":
                work[s["as_name"]] = d.dt.year.astype("Int64").astype(str) + " Q" + d.dt.quarter.astype("Int64").astype(str)
            elif part == "month":
                work[s["as_name"]] = d.dt.strftime("%Y-%m")
            else:
                raise PlanError(f"Step {i}: date_part needs 'year', 'quarter' or 'month'.")
        elif op == "calculate":
            _need(work, s["column"], i)
            a = _as_number(work[s["column"]])
            other = s.get("value", "")
            if other in work.columns:
                b = _as_number(work[other])
            else:
                bnum = sr.clean_amount(pd.Series([other])).iloc[0]
                if pd.isna(bnum):
                    raise PlanError(f"Step {i}: '{other}' is neither a column nor a number.")
                b = pd.Series(bnum, index=work.index)
            oper = s.get("operator")
            if oper == "add":
                r = a + b
            elif oper == "subtract":
                r = a - b
            elif oper == "multiply":
                r = a * b
            elif oper in ("divide", "percent_of"):
                r = a / b.where(b != 0)
                if oper == "percent_of":
                    r = r * 100
                zero = int((b == 0).sum())
                if zero:
                    caveats.append(f"Step {i}: {zero:,} rows divide by zero and are left blank.")
            else:
                raise PlanError(f"Step {i}: unknown calculation '{oper}'.")
            work[s["as_name"]] = r.round(2)
        elif op == "bucket":
            _need(work, s["column"], i)
            edges, labels = s.get("edges") or [], s.get("labels") or []
            if len(labels) != len(edges) - 1:
                raise PlanError(f"Step {i}: a bucket needs one fewer label than edges.")
            work[s["as_name"]] = pd.cut(_as_number(work[s["column"]]), bins=edges, labels=labels,
                                        include_lowest=True).astype(str).replace("nan", "")
        elif op == "group_summarize":
            by = s.get("columns") or []
            for c in by:
                _need(work, c, i)
            g = work.copy()
            aggs = {}
            for m in s.get("metrics") or []:
                src = m["column"]
                if m["agg"] == "count" and src not in g.columns:
                    src = by[0] if by else g.columns[0]
                _need(g, src, i)
                tmp = f"__{m['as_name']}"
                g[tmp] = g[src] if m["agg"] in ("count", "nunique") else _as_number(g[src])
                aggs[m["as_name"]] = (tmp, m["agg"])
            if not aggs:
                aggs = {"Rows": (by[0] if by else g.columns[0], "count")}
            if by:
                work = g.groupby(by, dropna=False).agg(**aggs).reset_index()
            else:
                work = pd.DataFrame({k: [g[c].agg(a)] for k, (c, a) in aggs.items()})
            for name, (_, agg) in aggs.items():
                if agg in ("sum", "mean"):
                    work[name] = work[name].round(2)
        elif op == "add_share":
            _need(work, s["column"], i)
            v = _as_number(work[s["column"]])
            total = v.sum()
            work[s["as_name"]] = (v / total * 100).round(1) if total else 0.0
        elif op == "sort":
            _need(work, s["column"], i)
            key = work[s["column"]]
            if not pd.api.types.is_numeric_dtype(key):
                if _is_dateish(key):
                    key = _as_date(key)
                elif sr._num_frac(key) >= 0.8:
                    key = _as_number(key)
            work = work.loc[key.sort_values(ascending=not s.get("descending"), na_position="last").index]
        elif op == "top_n":
            work = work.head(max(int(s.get("n") or 0), 0))
        elif op == "select_columns":
            for c in s.get("columns") or []:
                _need(work, c, i)
            work = work[list(s["columns"])]
        elif op == "rename_columns":
            work = work.rename(columns=dict(zip(s.get("columns") or [], s.get("values") or [])))
        else:
            raise PlanError(f"Step {i}: '{op}' is not an available operation.")
        if op == "filter" and before and not len(work):
            caveats.append(f"Step {i} left no rows — check the condition.")
    return work.reset_index(drop=True), caveats


# --------------------------------------------------------------------------
# Excel export
# --------------------------------------------------------------------------
def _sheet(title):
    for bad in '[]:*?/\\':
        title = title.replace(bad, " ")
    return (title.strip() or "Result")[:31]


def _style_header(ws):
    from openpyxl.styles import Font, PatternFill
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=sr.NAVY)
    for col in ws.columns:
        width = max(len(str(col[0].value or "")), *(len(str(c.value or "")) for c in col[1:200])) + 2
        ws.column_dimensions[col[0].column_letter].width = min(max(width, 10), 70)
    ws.freeze_panes = "A2"


def excel_bytes(history: list, plan: dict, result: pd.DataFrame, caveats: list, summary: str = "",
                source_names=(), health: list | None = None) -> bytes:
    """Final product: Answer (summary, request history, assumptions, steps, caveats), the result
    table, and the data review."""
    from openpyxl.styles import Alignment, Font, PatternFill
    if isinstance(history, str):
        history = [history]
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xl:
        pd.DataFrame().to_excel(xl, sheet_name="Answer", index=False)
        ws = xl.sheets["Answer"]
        ws.column_dimensions["A"].width = 24
        ws.column_dimensions["B"].width = 100
        rows = [("Answer", summary or "(summary not generated)"),
                ("As understood", plan.get("restated_objective", "")),
                ("Files", ", ".join(source_names))]
        rows += [(f"Objective {n}", h) for n, h in enumerate(history, 1)]
        rows += [("Assumption", a) for a in plan.get("assumptions", [])]
        rows += [(f"Step {n}", t.replace("**", "")) for n, t in enumerate(describe_steps(plan), 1)]
        rows += [("Caveat", c) for c in caveats]
        rows += [("Not covered", c) for c in plan.get("cannot_do", [])]
        rows += [("Method", "Steps were planned with AI from the objectives and the column list, approved, "
                            "then calculated with ordinary formulas from the original files. No AI produced "
                            "any number. AI-assisted draft — verify before relying on it.")]
        for r, (k, v) in enumerate(rows, 1):
            a, b = ws.cell(row=r, column=1, value=k), ws.cell(row=r, column=2, value=v)
            a.font = Font(bold=True, color="FFFFFF" if r == 1 else sr.NAVY)
            if r == 1:
                a.fill = PatternFill("solid", fgColor=sr.NAVY)
            b.alignment = Alignment(wrap_text=True, vertical="top")
        name = _sheet(plan.get("result_title") or "Result")
        if name == "Answer":
            name = "Result"
        result.to_excel(xl, sheet_name=name, index=False)
        _style_header(xl.sheets[name])
        if health:
            pd.DataFrame(health)[["priority", "file", "column", "issue", "fix", "rows_affected"]].rename(
                columns=str.title).to_excel(xl, sheet_name="Data review", index=False)
            _style_header(xl.sheets["Data review"])
    return buf.getvalue()
